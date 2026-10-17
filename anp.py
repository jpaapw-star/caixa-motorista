"""Robô da ANP: baixa a pesquisa semanal de preços dos combustíveis e grava dados/anp.json
com o preço médio de cada estado e de cada cidade pesquisada.

Roda sozinho no GitHub (veja .github/workflows/publicar.yml), duas vezes por semana.
Uso manual: python anp.py              (baixa a semana mais nova do site da ANP)
            python anp.py arquivo.xlsx (lê uma planilha já baixada)

Fonte: ANP, Levantamento de Preços de Combustíveis, "Preços médios semanais: Brasil, regiões, estados e municípios".
"""
import io
import json
import os
import re
import sys
import time
import unicodedata
import urllib.request
from datetime import date, datetime, timedelta, timezone
from urllib.parse import urljoin

from openpyxl import load_workbook

PAGINA = ("https://www.gov.br/anp/pt-br/assuntos/precos-e-defesa-da-concorrencia/precos/"
          "levantamento-de-precos-de-combustiveis-ultimas-semanas-pesquisadas")
PASTA_ARQUIVOS = "https://www.gov.br/anp/pt-br/assuntos/precos-e-defesa-da-concorrencia/precos/arquivos-lpc/"
SAIDA = os.path.join(os.path.dirname(os.path.abspath(__file__)), "dados", "anp.json")
CABECALHOS = {"User-Agent": "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/126.0 Safari/537.36"}

UFS = {"AC": "Acre", "AL": "Alagoas", "AP": "Amapá", "AM": "Amazonas", "BA": "Bahia", "CE": "Ceará",
       "DF": "Distrito Federal", "ES": "Espírito Santo", "GO": "Goiás", "MA": "Maranhão", "MT": "Mato Grosso",
       "MS": "Mato Grosso do Sul", "MG": "Minas Gerais", "PA": "Pará", "PB": "Paraíba", "PR": "Paraná",
       "PE": "Pernambuco", "PI": "Piauí", "RJ": "Rio de Janeiro", "RN": "Rio Grande do Norte",
       "RS": "Rio Grande do Sul", "RO": "Rondônia", "RR": "Roraima", "SC": "Santa Catarina", "SP": "São Paulo",
       "SE": "Sergipe", "TO": "Tocantins"}
COMBS = ["Gasolina", "Etanol", "Diesel", "GNV"]


def sa(t):
    """Sem acento, maiúsculas e espaços simples: 'São  Paulo' -> 'SAO PAULO'."""
    t = unicodedata.normalize("NFD", str(t if t is not None else "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", t).strip().upper()


UF_POR_NOME = {sa(n): uf for uf, n in UFS.items()}


def uf_de(valor):
    v = sa(valor)
    if v in UFS:
        return v
    return UF_POR_NOME.get(v)


def produto(valor):
    """Nome da ANP -> nome do app. Diesel comum só entra se não houver S10."""
    p = sa(valor)
    if p.startswith("GASOLINA") and "ADITIVADA" not in p and "AVIACAO" not in p:
        return "Gasolina"
    if p.startswith("ETANOL"):
        return "Etanol"
    if "DIESEL" in p:
        return "Diesel" if "S10" in p.replace(" ", "").replace("-", "") else "_diesel_comum"
    if p.startswith("GNV") or p == "GAS NATURAL VEICULAR":
        return "GNV"
    return None


def numero(v):
    if v is None or v == "" or v == "-":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    t = str(v).strip().replace("R$", "").replace(" ", "")
    if "," in t:
        t = t.replace(".", "").replace(",", ".")
    try:
        return float(t)
    except ValueError:
        return None


def nome_bonito(t):
    pequenas = {"de", "da", "do", "das", "dos", "e"}
    partes = str(t).strip().lower().split()
    return " ".join(p if (i and p in pequenas) else p[:1].upper() + p[1:] for i, p in enumerate(partes))


def data_iso(v):
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    m = re.match(r"(\d{1,2})/(\d{1,2})/(\d{4})", str(v or ""))
    return "%s-%02d-%02d" % (m.group(3), int(m.group(2)), int(m.group(1))) if m else None


def achar_colunas(cel):
    """Acha as colunas pelo nome do cabeçalho (sem depender da ordem nem da linha)."""
    def col(*comeca):
        for op in comeca:
            for j, c in enumerate(cel):
                if c == op or c.startswith(op):
                    return j
        return None
    return {
        "produto": col("PRODUTO"),
        "estado": col("ESTADO", "UF"),
        "municipio": col("MUNICIPIO", "CAPITAL", "CIDADE"),
        "regiao": col("REGIAO"),
        "medio": col("PRECO MEDIO REVENDA", "PRECO MEDIO"),
        "minimo": col("PRECO MINIMO REVENDA", "PRECO MINIMO"),
        "maximo": col("PRECO MAXIMO REVENDA", "PRECO MAXIMO"),
        "postos": col("NUMERO DE POSTOS", "N DE POSTOS", "QTD DE POSTOS"),
        "inicio": col("DATA INICIAL"),
        "fim": col("DATA FINAL"),
    }


def ler_planilha(conteudo):
    wb = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    estados, municipios, brasil = {}, {}, {}
    semana = {}
    for ws in wb.worksheets:
        linhas = ws.iter_rows(values_only=True)
        cols = None
        for i, row in enumerate(linhas):
            if cols is None:
                if i > 60:
                    break
                cel = [sa(c) for c in row]
                if "PRODUTO" in cel and any(c.startswith("PRECO MEDIO") for c in cel):
                    cols = achar_colunas(cel)
                    if cols["produto"] is None or cols["medio"] is None:
                        cols = None
                        break
                continue
            pega = lambda k: row[cols[k]] if cols[k] is not None and cols[k] < len(row) else None
            comb = produto(pega("produto"))
            medio = numero(pega("medio"))
            if not comb or not medio or medio <= 0:
                continue
            preco = {"medio": round(medio, 3)}
            for k in ("minimo", "maximo"):
                v = numero(pega(k))
                if v:
                    preco[k[:3]] = round(v, 3)
            n = numero(pega("postos"))
            if n:
                preco["postos"] = int(n)
            if not semana and pega("inicio") and pega("fim"):
                semana = {"inicio": data_iso(pega("inicio")), "fim": data_iso(pega("fim"))}
            uf = uf_de(pega("estado")) if cols["estado"] is not None else None
            if cols["municipio"] is not None and uf and pega("municipio"):
                nome = nome_bonito(pega("municipio"))
                chave = uf + "|" + sa(nome).lower()
                lugar = municipios.setdefault(chave, {"nome": nome, "uf": uf, "precos": {}})
            elif cols["municipio"] is None and uf:
                lugar = estados.setdefault(uf, {"nome": UFS[uf], "precos": {}})
            elif cols["municipio"] is None and cols["estado"] is None and cols["regiao"] is None:
                lugar = brasil.setdefault("BR", {"nome": "Brasil", "precos": {}})
            else:
                continue
            lugar["precos"][comb] = preco
    # diesel comum só quando não houver S10
    for lugar in list(estados.values()) + list(municipios.values()) + list(brasil.values()):
        p = lugar["precos"]
        comum = p.pop("_diesel_comum", None)
        if comum and "Diesel" not in p:
            p["Diesel"] = comum
    return estados, municipios, brasil.get("BR"), semana


def baixar(url, tentativas=3):
    for n in range(tentativas):
        try:
            with urllib.request.urlopen(urllib.request.Request(url, headers=CABECALHOS), timeout=90) as r:
                return r.read()
        except Exception as e:  # noqa: BLE001
            if n == tentativas - 1:
                raise
            print("  falhou (%s), tentando de novo…" % e)
            time.sleep(10 * (n + 1))


def fim_da_semana(url):
    datas = re.findall(r"(\d{4})-(\d{1,2})-(\d{1,2})", url)
    return date(*map(int, datas[-1])) if datas else date.min


def achar_arquivo():
    """Link da planilha mais nova na página da ANP; se a página mudar, tenta o nome padrão das últimas semanas."""
    candidatos = []
    try:
        html = baixar(PAGINA).decode("utf-8", "ignore")
        for href in re.findall(r'href="([^"]*resumo_semanal_lpc[^"]*\.xlsx)"', html, flags=re.I):
            candidatos.append(urljoin(PAGINA, href))
    except Exception as e:  # noqa: BLE001
        print("Não consegui abrir a página da ANP (%s). Tentando pelo nome do arquivo." % e)
    if candidatos:
        return sorted(set(candidatos), key=fim_da_semana)[-1]
    # a semana da ANP vai de domingo a sábado
    hoje = date.today()
    sabado = hoje - timedelta(days=(hoje.weekday() - 5) % 7 or 7)
    for k in range(5):
        fim = sabado - timedelta(weeks=k)
        ini = fim - timedelta(days=6)
        url = "%s%d/resumo_semanal_lpc_%s_%s.xlsx" % (PASTA_ARQUIVOS, fim.year, ini.isoformat(), fim.isoformat())
        try:
            req = urllib.request.Request(url, headers=CABECALHOS, method="HEAD")
            with urllib.request.urlopen(req, timeout=30) as r:
                if r.status == 200:
                    return url
        except Exception:  # noqa: BLE001
            continue
    raise SystemExit("Não achei a planilha semanal da ANP. O site pode ter mudado.")


def resumo(lugares):
    return {k: {c: v["precos"][c]["medio"] for c in COMBS if c in v["precos"]} for k, v in lugares.items()}


def main():
    if len(sys.argv) > 1:
        arquivo = sys.argv[1]
        with open(arquivo, "rb") as f:
            conteudo = f.read()
    else:
        arquivo = achar_arquivo()
        print("Baixando", arquivo)
        conteudo = baixar(arquivo)
    estados, municipios, brasil, semana = ler_planilha(conteudo)
    if not semana:
        datas = re.findall(r"(\d{4})-(\d{1,2})-(\d{1,2})", os.path.basename(arquivo))
        if len(datas) >= 2:
            semana = {"inicio": date(*map(int, datas[-2])).isoformat(), "fim": date(*map(int, datas[-1])).isoformat()}
    print("Lidos: %d estados e %d cidades. Semana: %s" % (len(estados), len(municipios), semana))
    if len(estados) < 20 or len(municipios) < 50:
        raise SystemExit("A planilha veio diferente do esperado (poucos estados ou cidades). Nada foi gravado.")

    novo = {
        "fonte": "ANP - Levantamento de Preços de Combustíveis (preços médios semanais)",
        "arquivo": os.path.basename(arquivo),
        "semana": semana,
        "gerado": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
        "brasil": brasil["precos"] if brasil else None,
        "estados": estados,
        "municipios": municipios,
    }
    # guarda a média da semana anterior para mostrar se subiu ou caiu
    antigo = None
    if os.path.exists(SAIDA):
        with open(SAIDA, encoding="utf-8") as f:
            antigo = json.load(f)
    if antigo and antigo.get("semana") == semana:
        if antigo.get("estados") == estados and antigo.get("municipios") == municipios:
            print("Semana já gravada, nada mudou.")
            return
        novo["anterior"] = antigo.get("anterior")
    elif antigo and antigo.get("estados"):
        novo["anterior"] = {"semana": antigo.get("semana"), "estados": resumo(antigo["estados"]),
                            "municipios": resumo(antigo.get("municipios") or {})}
    os.makedirs(os.path.dirname(SAIDA), exist_ok=True)
    with open(SAIDA, "w", encoding="utf-8") as f:
        json.dump(novo, f, ensure_ascii=False, separators=(",", ":"))
    print("Gravado em", SAIDA)


if __name__ == "__main__":
    main()
